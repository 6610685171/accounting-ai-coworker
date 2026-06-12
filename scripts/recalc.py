# ========= Copyright 2025-2026 @ Eigent.ai All Rights Reserved. =========
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
# ========= Copyright 2025-2026 @ Eigent.ai All Rights Reserved. =========

import openpyxl
import sys
import os

def check_excel_errors(file_path):
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return False

    print(f"Checking for errors in {file_path}...")
    try:
        # Note: openpyxl doesn't evaluate formulas, so it won't see #REF! 
        # unless the file was saved by Excel/LibreOffice with those values.
        # However, we can check for common formula issues if we were to evaluate them.
        
        wb = openpyxl.load_workbook(file_path, data_only=True)
        errors_found = 0
        
        error_values = ["#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NULL!", "#NUM!"]
        
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value in error_values:
                        print(f"  Error {cell.value} found in {sheet_name}!{cell.coordinate}")
                        errors_found += 1
        
        if errors_found == 0:
            print("No formula errors detected (in static values).")
            return True
        else:
            print(f"Total errors found: {errors_found}")
            return False
            
    except Exception as e:
        print(f"Error reading file: {e}")
        return False

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python recalc.py <file_path>")
    else:
        check_excel_errors(sys.argv[1])
